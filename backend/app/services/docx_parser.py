import html
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import bleach
from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from PIL import Image

from app.schemas.article import ArticleDocument
from app.services.errors import ServiceError
from app.services.filenames import slugify_filename, unique_filename

ALLOWED_TAGS = ["strong", "em", "a", "br"]
ALLOWED_ATTRS = {"a": ["href"]}
MAX_ARCHIVE_ENTRIES = 10_000
MAX_COMPRESSION_RATIO = 200
METADATA_PATTERN = re.compile(
    r"(?:^|\s+)(Meta Title|Meta Description|Excerpt|URL Slug)\s*:\s*"
    r"(.*?)(?=\s+(?:Meta Title|Meta Description|Excerpt|URL Slug)\s*:|$)",
    re.IGNORECASE,
)


@dataclass
class ExtractedImage:
    key: str
    original_filename: str
    filename: str
    mime_type: str
    data: bytes
    caption: str | None = None
    figure_number: int | None = None


@dataclass
class ParsedDocx:
    article: ArticleDocument
    images: list[ExtractedImage]


def validate_docx_archive(data: bytes, max_uncompressed_bytes: int) -> None:
    """Reject malformed or dangerous OOXML archives before XML parsing."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ServiceError("DOCX_SECURITY_ERROR", "The DOCX contains too many files.")
            total = 0
            for entry in entries:
                path = Path(entry.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ServiceError("DOCX_SECURITY_ERROR", "The DOCX contains an unsafe path.")
                total += entry.file_size
                if total > max_uncompressed_bytes:
                    raise ServiceError(
                        "DOCX_SECURITY_ERROR", "The expanded DOCX exceeds the safety limit.", 413
                    )
                if (
                    entry.file_size > 1_000_000
                    and entry.file_size > max(entry.compress_size, 1) * MAX_COMPRESSION_RATIO
                ):
                    raise ServiceError(
                        "DOCX_SECURITY_ERROR", "The DOCX contains unsafe compressed content."
                    )
            names = archive.namelist()
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ServiceError("DOCX_PARSE_ERROR", "The uploaded file is not a valid DOCX.")
    except ServiceError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise ServiceError("DOCX_PARSE_ERROR", "The uploaded file is not a valid DOCX.") from exc


def _iter_blocks(parent: DocumentType | _Cell):
    parent_element = parent.element.body if isinstance(parent, DocumentType) else parent._tc
    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _inline_html(paragraph: Paragraph) -> str:
    parts: list[str] = []
    for child in paragraph._p.iterchildren():
        if child.tag == qn("w:r"):
            texts = child.xpath(".//w:t/text()")
            breaks = len(child.xpath(".//w:br"))
            value = html.escape("".join(texts)) + "<br>" * breaks
            if not value:
                continue
            props = child.find(qn("w:rPr"))
            if props is not None and props.find(qn("w:i")) is not None:
                value = f"<em>{value}</em>"
            if props is not None and props.find(qn("w:b")) is not None:
                value = f"<strong>{value}</strong>"
            parts.append(value)
        elif child.tag == qn("w:hyperlink"):
            rel_id = child.get(qn("r:id"))
            href = paragraph.part.rels[rel_id].target_ref if rel_id in paragraph.part.rels else ""
            label = html.escape("".join(child.xpath(".//w:t/text()")))
            parts.append(f'<a href="{html.escape(href, quote=True)}">{label}</a>')
    return bleach.clean("".join(parts), tags=ALLOWED_TAGS, attributes=ALLOWED_ATTRS, strip=True)


def _style_level(paragraph: Paragraph) -> int | None:
    name = (paragraph.style.name if paragraph.style else "").lower()
    match = re.match(r"heading\s*(\d+)", name)
    return int(match.group(1)) if match else None


def _is_list(paragraph: Paragraph) -> tuple[bool, bool]:
    style = (paragraph.style.name if paragraph.style else "").lower()
    if style.startswith("list"):
        return True, "number" in style
    ppr = paragraph._p.pPr
    if ppr is None or ppr.numPr is None:
        return False, False
    return True, "number" in style


def _image_rel_ids(paragraph: Paragraph) -> list[str]:
    return paragraph._p.xpath(".//a:blip/@r:embed")


def _metadata_fields(text: str) -> dict[str, str]:
    keys = {
        "meta title": "meta_title",
        "meta description": "meta_description",
        "excerpt": "excerpt",
        "url slug": "slug",
    }
    if not re.match(r"^\s*(?:Meta Title|Meta Description|Excerpt|URL Slug)\s*:", text, re.I):
        return {}
    return {
        keys[match.group(1).lower()]: match.group(2).strip()
        for match in METADATA_PATTERN.finditer(text)
    }


def parse_docx(data: bytes) -> ParsedDocx:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ServiceError("DOCX_PARSE_ERROR", "The DOCX file could not be opened.") from exc

    metadata: dict[str, str] = {}
    title = ""
    blocks: list[dict] = []
    images: list[ExtractedImage] = []
    used_names: set[str] = set()

    for item in _iter_blocks(document):
        if isinstance(item, Table):
            rows = [
                [
                    {
                        "html": bleach.clean(
                            _inline_html(cell.paragraphs[0]) if cell.paragraphs else "",
                            tags=ALLOWED_TAGS,
                            attributes=ALLOWED_ATTRS,
                        )
                    }
                    for cell in row.cells
                ]
                for row in item.rows
            ]
            blocks.append({"type": "table", "rows": rows})
            continue

        text = item.text.strip()
        ids = _image_rel_ids(item)
        for rel_id in ids:
            part = document.part.related_parts[rel_id]
            original = Path(str(part.partname)).name
            raw = part.blob
            try:
                with Image.open(io.BytesIO(raw)) as image:
                    image.verify()
                    mime = Image.MIME.get(image.format or "", part.content_type)
            except Exception as exc:
                raise ServiceError(
                    "IMAGE_EXTRACTION_ERROR", f"Embedded image {original} is invalid."
                ) from exc
            key = f"image-{len(images) + 1}"
            image = ExtractedImage(key, original, original, mime, raw)
            images.append(image)
            blocks.append({"type": "image", "image_id": key, "source": key})

        if ids and not text:
            continue
        caption_match = re.match(r"^Figure\s+(\d+)\s*[:.\-]?", text, re.I)
        style_name = (item.style.name or "").lower() if item.style else ""
        is_caption = bool(caption_match) or "caption" in style_name
        if is_caption and text and blocks and blocks[-1]["type"] == "image":
            current = images[-1]
            current.caption = text
            current.figure_number = int(caption_match.group(1)) if caption_match else None
            current.filename = unique_filename(
                slugify_filename(text, Path(current.original_filename).suffix, current.key),
                used_names,
            )
            blocks[-1].update(
                caption=text, figure_number=current.figure_number, filename=current.filename
            )
            continue

        parsed_metadata = _metadata_fields(text)
        if parsed_metadata:
            metadata.update(parsed_metadata)
        else:
            if not text:
                continue
            level = _style_level(item)
            if level == 1 and not title:
                title = text
            elif level in (2, 3):
                blocks.append({"type": "heading", "level": level, "text": text})
            else:
                listed, ordered = _is_list(item)
                value = _inline_html(item)
                if (
                    listed
                    and blocks
                    and blocks[-1]["type"] == "list"
                    and blocks[-1]["ordered"] == ordered
                ):
                    blocks[-1]["items"].append({"html": value})
                elif listed:
                    blocks.append({"type": "list", "ordered": ordered, "items": [{"html": value}]})
                else:
                    blocks.append({"type": "paragraph", "html": value})

    for image in images:
        if image.filename == image.original_filename:
            image.filename = unique_filename(image.original_filename, used_names)
            for block in blocks:
                if block.get("image_id") == image.key:
                    block["filename"] = image.filename
                    break
    if not title:
        raise ServiceError("DOCX_PARSE_ERROR", "The document must contain an H1 article title.")
    return ParsedDocx(
        ArticleDocument(title=title, blocks=blocks, **metadata),
        images,
    )
