import io
import zipfile

import pytest
from docx import Document
from PIL import Image

from app.services.docx_parser import parse_docx, validate_docx_archive
from app.services.errors import ServiceError
from app.services.filenames import slugify_filename, unique_filename
from app.services.renderer import render_article


def sample_docx() -> bytes:
    document = Document()
    document.add_paragraph("Meta Title: Exact SEO title")
    document.add_paragraph("Meta Description: Exact SEO description")
    document.add_paragraph("URL Slug: exact-source-slug")
    document.add_heading("Exact Article Title", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Bold").bold = True
    paragraph.add_run(" and ")
    paragraph.add_run("italic").italic = True
    document.add_heading("Section", level=2)
    document.add_paragraph("Bullet one", style="List Bullet")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    image = Image.new("RGB", (20, 20), "red")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")
    document.add_paragraph().add_run().add_picture(io.BytesIO(image_bytes.getvalue()))
    document.add_paragraph("Figure 1: Mine site's overview. Image credit: Publisher.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_parser_preserves_metadata_order_formatting_table_and_image() -> None:
    parsed = parse_docx(sample_docx())
    assert parsed.article.title == "Exact Article Title"
    assert parsed.article.meta_title == "Exact SEO title"
    assert parsed.article.meta_description == "Exact SEO description"
    assert parsed.article.slug == "exact-source-slug"
    assert [block.type for block in parsed.article.blocks] == [
        "paragraph",
        "heading",
        "list",
        "table",
        "image",
    ]
    assert parsed.article.blocks[0].html == "<strong>Bold</strong> and <em>italic</em>"
    assert parsed.images[0].filename == "mine-sites-overview.png"
    assert parsed.images[0].caption == "Figure 1: Mine site's overview. Image credit: Publisher."


def test_parser_extracts_multiple_metadata_fields_from_one_paragraph() -> None:
    document = Document()
    document.add_paragraph(
        "Meta Title: Capricorn Metals Karlawinda Output Strategy Update "
        "Meta Description: Capricorn Metals completes a major processing upgrade at "
        "Karlawinda. See what this output strategy shift means for CMM on Colitco. "
        "URL Slug: capricorn-metals-karlawinda-output-strategy"
    )
    document.add_heading("Capricorn Metals Article", level=1)
    document.add_paragraph("Exact body content.")
    output = io.BytesIO()
    document.save(output)

    article = parse_docx(output.getvalue()).article

    assert article.meta_title == "Capricorn Metals Karlawinda Output Strategy Update"
    assert article.meta_description == (
        "Capricorn Metals completes a major processing upgrade at Karlawinda. "
        "See what this output strategy shift means for CMM on Colitco."
    )
    assert article.slug == "capricorn-metals-karlawinda-output-strategy"


def test_filename_normalization_and_duplicates() -> None:
    filename = slugify_filename("Figure 2: IGO’s Nova—site! Image credit: IGO.", ".JPG", "image")
    used: set[str] = set()
    assert filename == "igos-novasite.jpg"
    assert unique_filename(filename, used) == "igos-novasite.jpg"
    assert unique_filename(filename, used) == "igos-novasite-2.jpg"


def test_renderer_centers_image_and_keeps_caption() -> None:
    article = parse_docx(sample_docx()).article
    rendered = render_article(article, {"image-1": "https://example.test/image.png"})
    assert 'class="wp-block-image aligncenter"' in rendered
    assert "Figure 1: Mine site&#x27;s overview. Image credit: Publisher." in rendered
    assert "<h1>" not in rendered


def test_docx_archive_validation_accepts_a_real_document() -> None:
    validate_docx_archive(sample_docx(), 10 * 1024 * 1024)


def test_docx_archive_validation_rejects_path_traversal() -> None:
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "document")
        archive.writestr("../outside.txt", "unsafe")

    with pytest.raises(ServiceError, match="unsafe path") as raised:
        validate_docx_archive(payload.getvalue(), 10 * 1024 * 1024)

    assert raised.value.code == "DOCX_SECURITY_ERROR"


def test_docx_archive_validation_enforces_expanded_size() -> None:
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "x" * 2048)

    with pytest.raises(ServiceError) as raised:
        validate_docx_archive(payload.getvalue(), 1024)

    assert raised.value.code == "DOCX_SECURITY_ERROR"
